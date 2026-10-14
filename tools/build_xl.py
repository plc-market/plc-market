import openpyxl, json, re, statistics as st, copy
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
SRC='xl/src.xlsx'; OUT='xl/Siemens_Otomasyon_Fiyat_Listesi_Site.xlsx'
paths=json.load(open('paths.json')); VALID={p[3] for p in paths}
wb=openpyxl.load_workbook(SRC)
toc={}; up={}
for r in wb['İçindekiler'].iter_rows(min_row=5,max_row=53,values_only=True): toc[r[1]]=r[2]; up[r[1]]=r[0]
def P(a,b,c):
    p=f"{a} > {b} > {c}"; assert p in VALID,p; return p
K="Kontrol sistemleri";S="Sürücü teknolojisi";FI="Frekans invertörleri";SV="Servo sürücü sistemleri"
def et200sp(d):
    D=d.upper()
    if 'CPU' in D: return 'Standart CPU\'lar'
    if re.search(r'BASEUNIT|BASE UNIT|\bBU\d|\bBU ',D): return 'Baz üniteler'
    if re.search(r'IM ?155|ARAYÜZ|INTERFACE',D): return 'Arayüz modülleri'
    if re.search(r'BUS ?ADAP|\bBA ',D): return 'Bus adaptörleri'
    if re.search(r'MOTOR ?START|MOTOR YOL',D): return 'ET 200SP motor yol vericiler'
    if re.search(r'SIWAREX|TARTIM|WEIGH',D): return 'ET 200SP tartım modülleri'
    if re.search(r'\bCM\b|\bCP\b|HABERLEŞME|COMMUNICATION|AS-I|IO-LINK|SERIAL|SERİ',D): return 'ET 200SP haberleşme modülleri'
    if re.search(r'\bAI\b|\bAQ\b|ANALOG',D): return 'Analog G/Ç modülleri'
    if re.search(r'\bDI\b|\bDQ\b|DİJİTAL|DIGITAL|\bF-DI|\bF-DQ|\bRQ\b',D): return 'Dijital G/Ç modülleri'
    if re.search(r'\bPM\b|POWER|GÜÇ|\bPS\b',D): return 'ET 200SP güç kaynakları'
    return 'ET 200SP hafıza kartları / aksesuarlar'
def cat(sh,ara,desc,sku):
    a=(ara or ''); A=a.lower(); d=desc or ''
    acc=('aksesuar' in A or 'opsiyon' in A)
    if sh=='SINAMICS V20': return P(S,FI,'Frekans invertörü aksesuarları' if acc else 'SINAMICS V20 frekans invertörleri')
    if sh=='SINAMICS G120C': return P(S,FI,'Frekans invertörü aksesuarları' if acc else 'SINAMICS G120C frekans invertörleri')
    if sh=='SINAMICS G120X': return P(S,FI,'SINAMICS G120X frekans invertörleri')
    if sh=='SINAMICS G120': return P(S,FI,'Frekans invertörü aksesuarları' if acc else 'SINAMICS G120 modüler invertörler')
    if sh in('SINAMICS G220','SINAMICS G220 Clean Power'): return P(S,FI,'Frekans invertörü aksesuarları' if acc else 'SINAMICS G220 frekans invertörleri')
    if sh=='SINAMICS G130': return P(S,FI,'Frekans invertörü aksesuarları' if acc else 'SINAMICS G130 frekans invertörleri')
    if sh=='SINAMICS DCM': return P(S,FI,'Frekans invertörü aksesuarları' if acc else 'SINAMICS DCM DC sürücüler')
    if sh=='SIMATIC Micro Drive': return P(S,SV,'SIMATIC MICRO-DRIVE')
    if sh=='SINAMICS V90': return P(S,SV,'SINAMICS V90 servo sistemleri')
    if sh=='SINAMICS S200': return P(S,SV,'SINAMICS S200 servo setleri')
    if sh=='SINAMICS S210': return P(S,SV,'SINAMICS S210 servo setleri')
    if sh in('S110 ve S120 Kompakt','S120 Modüler'): return P(S,SV,'SINAMICS S110 / S120 sürücüler')
    if sh=='S120 Servo Motor Kabloları': return P(S,SV,'Servo motor kabloları')
    if sh=='SIMOTICS 1FK7': return P(S,'Motorlar','SIMOTICS 1FK7 servo motorlar')
    if sh=='SIMOTICS 1PH8': return P(S,'Motorlar','SIMOTICS 1PH8 asenkron motorlar')
    if sh=='SINUMERIK 808D': return P(K,'Hareket ve CNC kontrolörleri','SINUMERIK 808D aksesuar ve kabloları' if a in('Kablolar','El çarkı') else 'SINUMERIK 808D CNC sistemleri')
    R='RFID ve kod okuma'; SE='Sensör sistemleri'
    if sh in('RFID RF600','RFID RF200 ve RF300','MOBY D - RF300 ISO Tag'):
        if 'tag' in A: return P(SE,R,"RFID tag'ler")
        if 'terminal' in A: return P(SE,R,'RFID el terminalleri')
        if 'erişim' in A or 'access' in A: return P(SE,R,'RFID erişim kontrolü')
        return P(SE,R,'RFID okuyucular ve sistem bileşenleri')
    if sh=='RFID Arabirim Kabloları': return P(SE,R,'RFID haberleşme modülleri ve kabloları')
    if sh=='2D Karekod Okuma': return P(SE,R,'El tipi kod okuyucular' if 'el tipi' in A else 'Sabit kod okuyucular')
    G='Güç kaynakları'; GU='Anahtarlamalı güç kaynakları / UPS'
    if sh=='SITOP Güç Kaynakları':
        if 'logo! power' in A: return P(K,'LOGO! mini kontrolörler','LOGO! güç kaynağı')
        if 'dc/dc' in A: return P(G,GU,'DC/DC dönüştürücüler')
        if re.search(r'redundancy|selectivity|buffer|inrush|seçicilik',A): return P(G,GU,'Güç kaynağı ek modülleri')
        if 'akü' in A: return P(G,GU,'UPS akü modülleri')
        if 'ups' in A: return P(G,GU,'Kesintisiz güç kaynakları')
        if 'ip67' in A or 'psu100p' in A: return P(G,GU,'IP67 güç kaynakları, 24 VDC')
        if 'flexi' in A or 'dual' in A: return P(G,GU,'Ray tipi güç kaynakları, çok gerilimli')
        return P(G,GU,'Ray tipi güç kaynakları, 24 VDC')
    LG='LOGO! mini kontrolörler'
    if sh=='LOGO! 8.4':
        if 'basic' in A: return P(K,LG,'Ekransız lojik modüller' if str(sku).startswith('6ED1052-2') else 'Ekranlı lojik modüller')
        if 'dijital' in A: return P(K,LG,'Dijital genişleme modülleri')
        if 'analog' in A: return P(K,LG,'Analog genişleme modülleri')
        if 'knx' in A: return P(K,LG,'Haberleşme modülleri')
        if 'text' in A: return P(K,LG,'Metin ekranları')
        if 'soft' in d.lower(): return P(K,LG,'LOGO! programlama yazılımı')
        return P(K,LG,'LOGO! aksesuarları')
    S12='SIMATIC S7-1200'; GV='Güvenlik teknolojisi'; GK='Güvenlik kontrolörleri'
    if sh=='SIMATIC S7-1200':
        if 'fc cpu' in A or '1200f' in A: return P(GV,GK,'Failsafe SIMATIC S7-1200F')
        if 'kart' in A and 'haberleşme' not in A: return P(K,S12,'Sinyal kartları (boards)')
        if 'cpu' in A: return P(K,S12,'Merkezi işlem birimleri (CPU)')
        if 'dijital' in A: return P(K,S12,'Dijital giriş/çıkış modülleri')
        if 'analog' in A or 'enerji' in A: return P(K,S12,'Analog giriş/çıkış modülleri')
        if 'telekontrol' in A: return P(K,S12,'Telekontrol modülleri ve yazılımları')
        if 'haberleşme' in A: return P(K,S12,'Haberleşme işlemcileri')
        if 'güç' in A: return P(K,S12,'S7-1200 güç kaynakları')
        if 'yazılım' in A: return P(K,S12,'STEP 7 Basic programlama yazılımı')
        return P(K,S12,'Hafıza kartları / aksesuarlar')
    H='HMI paneller'
    if sh=='HMI Operatör Paneller':
        if 'unified comfort pro' in A or 'comfort pro' in A: return P(K,H,'SIMATIC HMI Unified Comfort PRO paneller')
        if 'opsiyon' in A or 'aksesuar' in A or 'ekipman' in A: return P(K,H,'HMI panel aksesuarları')
        if 'unified basic' in A: return P(K,H,'SIMATIC HMI Unified Basic paneller')
        if 'unified comfort' in A: return P(K,H,'SIMATIC HMI Unified Comfort paneller')
        if 'comfort' in A: return P(K,H,'SIMATIC HMI Comfort paneller')
        if 'basic' in A: return P(K,H,'SIMATIC HMI Basic paneller (2. nesil)')
        if 'projelendirme' in A: return P(K,H,'WinCC Comfort programlama yazılımı')
        if 'key' in A: return P(K,H,'SIMATIC Key paneller')
        if 'paslanmaz' in A or 'outdoor' in A: return P(K,H,'Paslanmaz ve dış ortam panelleri')
        if 'mobil' in A: return P(K,H,'SIMATIC HMI mobil paneller')
        return P(K,H,'HMI panel aksesuarları')
    Y='Otomasyon yazılımları'
    m={'TIA HMI Projelendirme':'WinCC projelendirme yazılımları','TIA Runtime Advanced':'WinCC Runtime Advanced','TIA Runtime Professional':'WinCC Runtime Professional',
       'TIA HMI Runtime Opsiyonları':'WinCC opsiyonları','WinCC Unified Runtime':'WinCC Unified PC Runtime','WinCC SCADA':'WinCC SCADA (V8)','WinCC SCADA Opsiyonları':'WinCC opsiyonları','Power Management':'Enerji yönetim yazılımları'}
    if sh in m: return P(K,Y,m[sh])
    HC='Hareket ve CNC kontrolörleri'
    if sh=='SIMOTION':
        if '1500t' in A: return P(K,HC,"S7-1500T teknoloji CPU'ları")
        if 'scout' in A: return P(K,HC,'SIMOTION mühendislik yazılımları')
        return P(K,HC,'SIMOTION D hareket kontrolörleri')
    if sh=='SIMATIC S7-400':
        T='SIMATIC S7-400'
        if 'cpu' in A or 'bundle' in A: return P(K,T,"S7-400 CPU'lar")
        if 'dijital' in A or 'analog' in A: return P(K,T,'S7-400 G/Ç modülleri')
        if 'fonksiyon' in A or 'haberleşme' in A: return P(K,T,'S7-400 haberleşme ve fonksiyon modülleri')
        return P(K,T,'S7-400 rack, güç kaynağı ve hafıza')
    if sh=='SIMATIC S7-1500':
        T='SIMATIC S7-1500'
        if 'failsafe' in A: return P(GV,GK,'Failsafe SIMATIC S7-1500 / S7-400')
        if 'cpu' in A: return P(K,T,"S7-1500 CPU'lar")
        if '200mp' in A: return P(K,T,'ET 200MP modülleri')
        if 'dijital' in A: return P(K,T,'S7-1500 dijital G/Ç modülleri')
        if 'analog' in A: return P(K,T,'S7-1500 analog G/Ç modülleri')
        if 'haberleşme' in A: return P(K,T,'S7-1500 haberleşme modülleri')
        if 'teknoloji' in A: return P(K,T,'S7-1500 teknoloji modülleri')
        return P(K,T,'S7-1500 aksesuarları')
    if sh=='ET200SP Open - Software Ctrl':
        if 'open controller' in A: return P(K,'SIMATIC ET 200SP','ET 200SP Open Controller')
        return P(K,Y,'Software Controller ve ODK')
    if sh=='SIMATIC Yazılım Paketleri':
        if 'safety' in A: return P(GV,GK,'Güvenlik kontrolörü programlama yazılımı')
        if 'manager' in A: return P(K,Y,'SIMATIC Manager yazılımları')
        return P(K,Y,'TIA Portal ve STEP 7')
    if sh=='SIMATIC SAFETY':
        if 'cpu' in A: return P(GV,GK,'Failsafe SIMATIC S7-1500 / S7-400')
        if 'yazılım' in A: return P(GV,GK,'Güvenlik kontrolörü programlama yazılımı')
        if '200sp' in A: return P(GV,GK,'Failsafe SIMATIC ET 200SP')
        return P(GV,GK,'Failsafe dağıtılmış G/Ç')
    if sh=='SIRIUS SAFETY':
        if 'as-i' in A: return P(K,'SIMATIC ET 200SP','ET 200SP haberleşme modülleri')
        if 'konnektör' in A: return P(GV,GK,'Güvenlik kontrolörü aksesuarları')
        if 'başlangıç' in A: return P(GV,GK,'Güvenlik kontrolörü programlama yazılımı')
        return P(GV,GK,'Güvenlik mini kontrolörleri 3SK2')
    if sh=='ET200 Dağıtılmış IO':
        E='SIMATIC ET 200 dağıtılmış G/Ç'
        if '200sp' in A: return P(K,'SIMATIC ET 200SP',et200sp(d))
        if '200al' in A: return P(K,E,'ET 200AL modülleri')
        if 'eco' in A: return P(K,E,'ET 200eco PN modülleri')
        if 'pro' in A: return P(K,E,'ET 200pro modülleri')
        if 'isp' in A: return P(K,E,'ET 200iSP modülleri')
        return P(K,E,'ET 200M Ex modülleri')
    if sh=='SIMATIC TOP Connect':
        T='SIMATIC TOP Connect'
        if 'yapılmış' in A: return P(K,T,'Hazır kablolu ön konnektörler')
        if 'ön konnektör' in A: return P(K,T,'Ön konnektör modülleri')
        if 'kablo' in A: return P(K,T,'TOP Connect bağlantı kabloları')
        return P(K,T,'TOP Connect bağlantı modülleri')
    HB='Endüstriyel haberleşme'; EE='Endüstriyel Ethernet'; PN='PROFINET'
    if sh=='SIMATIC NET':
        if 'profibus pc' in A: return P(HB,'PROFIBUS','PROFIBUS PC kartları')
        if 'profibus' in A: return P(HB,'PROFIBUS','PROFIBUS haberleşme ürünleri')
        if 'genel' in A: return P(HB,"Gateway'ler","PROFINET / PROFIBUS gateway'ler")
        if 'fiber' in A or 'dönüştürücü' in A: return P(HB,EE,'Fiber optik kablolar ve konnektörler')
        if 'konnektörlü ethernet' in A: return P(HB,PN,'RJ45 PROFINET patch kabloları')
        if 'kablo' in A and 'anten' not in A: return P(HB,PN,'PROFINET kabloları')
        if 'modem' in A: return P(HB,"Endüstriyel router'lar","Endüstriyel LTE router'lar")
        if 'ethernet pc' in A: return P(HB,EE,'Endüstriyel Ethernet PC kartları')
        if 'opc' in A or 'nms' in A or 'traffic' in A: return P(HB,EE,'OPC server ve ağ yönetim yazılımları')
        if 'lpe' in A: return P(K,'IoT bileşenleri',"IoT gateway'ler")
        if 'scalance s' in A: return P(HB,"Endüstriyel router'lar",'Endüstriyel güvenlik duvarları')
        if 'sinema' in A: return P(HB,"Endüstriyel router'lar",'Endüstriyel VPN hizmetleri')
        if 'rtu' in A: return P(HB,EE,'Uzaktan erişim ve RTU')
        if 'sfp' in A: return P(HB,EE,'SFP modülleri')
        if 'xm 400' in A or 'x500' in A: return P(HB,EE,"Yönetilebilir Gigabit switch'ler (XM400 / X500)")
        if re.search(r'x[bcr]-100',A): return P(HB,PN,"Yönetilmeyen PROFINET switch'ler")
        if 'anten' in A or 'rcoax' in A: return P(HB,'Endüstriyel kablosuz','Anten ve RCoax ürünleri')
        if 'kablosuz' in A: return P(HB,'Endüstriyel kablosuz','Kablosuz Ethernet')
        return P(HB,PN,"Yönetilebilir PROFINET switch'ler")
    IP="Endüstriyel PC'ler ve monitörler"
    if sh=='Programlama Cihazları': return P(K,IP,'Programlama cihazları (Field PG)')
    if sh=='Endüstriyel PC ve Monitörler':
        if 'wincc' in A: return P(K,IP,'PC için yazılım paketleri')
        if 'tablet' in A: return P(K,IP,"Endüstriyel tablet PC'ler")
        if re.search(r'ifp|flat panel|thin client|itc',A): return P(K,IP,'Endüstriyel monitörler ve thin client')
        if re.search(r'rack|rw-|rc-|rs-|547j|647e|847e',A): return P(K,IP,"Rack PC'ler")
        if re.search(r'panel pc|nanopanel|px-|377g',A): return P(K,IP,"Panel PC'ler")
        return P(K,IP,"Box PC'ler")
    return None   # SITRAIN, hizmetler
# ---- piyasa verisi
BR=55.2105
mk=json.load(open('mk.json'))
exec(open('ratios.py').read().split("for k,v in T.items(): ob(k,float(v))")[0].split("obs={}")[1].replace("def ob(k,v): obs.setdefault(k.replace(' ',''),[]).append(v)\nfor k,v in mk.items(): ob(k,v[0]/BR)",""),globals()) if False else None
src=open('ratios.py',encoding='utf-8').read()
T=dict(x.split() for x in re.search(r'T=dict\(x\.split\(\) for x in """(.*?)"""',src,re.S).group(1).split('\n'))
D=dict(x.split() for x in re.search(r'D=dict\(x\.split\(\) for x in """(.*?)"""',src,re.S).group(1).split('\n'))
MKT={}
for k,v in mk.items(): MKT[k]=(v[0]/BR,v[1])
for k in set(T)|set(D):
    if k in MKT: continue
    vals=[float(x[k]) for x in (T,D) if k in x]
    MKT[k]=(sum(vals)/len(vals),len(vals))
# aile oranları
LIST={}
for ws in wb.worksheets:
    if ws.title in('İçindekiler','Tüm Ürünler'): continue
    hdr=[c.value for c in ws[4]]; ci=hdr.index('Sipariş No'); pi=hdr.index('Liste Fiyatı (EUR)')
    for r in ws.iter_rows(min_row=5,values_only=True):
        if r[ci] and isinstance(r[pi],(int,float)) and r[pi]>0: LIST.setdefault(str(r[ci]).replace(' ',''),(r[pi],ws.title))
fr={};ur={}
for k,(m,n) in MKT.items():
    if k in LIST:
        l,sh=LIST[k]; fr.setdefault(sh,[]).append(m/l); ur.setdefault(up[sh],[]).append(m/l)
FR={k:st.median(v) for k,v in fr.items()}; UR={k:st.median(v) for k,v in ur.items()}
print('aile oranları',{k:(round(v,3),len(fr[k])) for k,v in FR.items()}); print('üst grup',{k:(round(v,3),len(ur[k])) for k,v in UR.items()})
DATE='11–14.10.2026'
def price(sku,lst,sh):
    k=sku.replace(' ','')
    if k in MKT: m,n=MKT[k]; return round(m,2),f"Piyasa ortalaması · {n} site · {DATE}"
    if not isinstance(lst,(int,float)) or lst<=0: return None,"Fiyat yok"
    if sh in('SITRAIN Eğitimleri','Sürdürülebilirlik Hizmetleri'): return round(lst,2),"Liste fiyatı (hizmet)"
    if sh in FR: r=FR[sh]; return round(lst*r,2),f"Tahmini · liste × %{r*100:.1f} (bu ailede {len(fr[sh])} ürünün piyasa ortalaması)"
    u=up[sh]; r=UR.get(u, st.median([x for v in ur.values() for x in v]))
    return round(lst*r,2),f"Tahmini · liste × %{r*100:.1f} ({u} piyasa örneklerinin medyanı)"
# ---- stiller
HF=PatternFill('solid',fgColor='FF0F4FA8'); HFont=Font(b=True,color='FFFFFFFF'); NEWF=PatternFill('solid',fgColor='FFEAF1FB')
NEW=['Sitede Aktif','Site Kategorisi','Piyasa Ortalaması (EUR)','Nihai Satış Fiyatı (EUR)','Fiyat Dayanağı','Stok Adedi']
# kategori sayfası
cs=wb.create_sheet('Site Kategorileri',1)
cs.append(['Site Kategorisi (Ana > Grup > Alt)','Ana kategori','Grup','Alt kategori']); 
for p in paths: cs.append([p[3],p[0],p[1],p[2]])
for c in cs[1]: c.fill=HF; c.font=HFont
cs.column_dimensions['A'].width=90; cs.column_dimensions['B'].width=28; cs.column_dimensions['C'].width=34; cs.column_dimensions['D'].width=44; cs.freeze_panes='A2'
NP=len(paths)+1
seen=set(); stats={'aktif':0,'pasif':0,'piyasa':0,'tahmini':0,'kategorisiz':0}
ROWS={}
for ws in wb.worksheets:
    if ws.title in('İçindekiler','Tüm Ürünler','Site Kategorileri'): continue
    hdr=[c.value for c in ws[4]]; n0=len(hdr)
    ix={h:i for i,h in enumerate(hdr)}
    for j,h in enumerate(NEW):
        c=ws.cell(4,n0+1+j,h); c.font=HFont; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
    wid=[12,58,16,16,44,12]
    for j,w in enumerate(wid): ws.column_dimensions[get_column_letter(n0+1+j)].width=w
    dv1=DataValidation(type='list',formula1='"Evet,Hayır"',allow_blank=True); dv2=DataValidation(type='list',formula1=f"='Site Kategorileri'!$A$2:$A${NP}",allow_blank=True,showErrorMessage=True,error="Site Kategorileri sekmesindeki yollardan birini seçin.",errorTitle="Geçersiz kategori")
    ws.add_data_validation(dv1); ws.add_data_validation(dv2)
    last=ws.max_row
    for r in range(5,last+1):
        sku=ws.cell(r,ix['Sipariş No']+1).value
        if not sku: continue
        sku=str(sku).strip()
        ara=ws.cell(r,ix['Ara Grup']+1).value; desc=ws.cell(r,ix['Açıklama']+1).value
        lst=ws.cell(r,ix['Liste Fiyatı (EUR)']+1).value
        opt=ws.cell(r,ix['Seçenek Kodu']+1).value if 'Seçenek Kodu' in ix else None
        c=cat(ws.title,ara,desc,sku)
        pr,basis=price(sku,lst,ws.title)
        key=sku.replace(' ','')
        act = c is not None and pr is not None and not opt and key not in seen
        if opt: basis="Seçenek kodu: ana ürüne eklenen ek ücret, tek başına satılmaz"
        elif key in seen and c and pr: basis+=" · aynı sipariş no başka sekmede aktif"
        if act: seen.add(key)
        stats['aktif' if act else 'pasif']+=1
        if c is None: stats['kategorisiz']+=1
        if basis.startswith('Piyasa'): stats['piyasa']+=1
        elif basis.startswith('Tahmini'): stats['tahmini']+=1
        vals=['Evet' if act else 'Hayır', c, pr if not opt else None, pr if not opt else None, basis, None]
        for j,v in enumerate(vals):
            cell=ws.cell(r,n0+1+j,v); cell.fill=NEWF
            if j in(2,3): cell.number_format='#,##0.00 €'
        dv1.add(ws.cell(r,n0+1)); dv2.add(ws.cell(r,n0+2))
        ROWS[(ws.title,r)]=vals
        ROWS.setdefault(('k',key,ws.title),vals)
    if ws.auto_filter.ref: ws.auto_filter.ref=f"A4:{get_column_letter(n0+len(NEW))}{last}"
# Tüm Ürünler: aynı sütunlar (referans kopyası)
ws=wb['Tüm Ürünler']; hdr=[c.value for c in ws[1]]; n0=len(hdr)
for j,h in enumerate(NEW):
    c=ws.cell(1,n0+1+j,h); c.font=HFont; c.fill=HF
    ws.column_dimensions[get_column_letter(n0+1+j)].width=[12,58,16,16,44,12][j]
inv={v:k for k,v in toc.items()}
# sekme sırasıyla eşleştir: aynı (sekme, sipariş no) sırası
order={}
for (k,*rest),v in [((a[0],)+tuple(a[1:]),b) for a,b in ROWS.items() if a[0]!='k']: pass
seq={}
for key,v in ROWS.items():
    if key[0]=='k': continue
    sh,r=key; seq.setdefault(sh,[]).append((r,v))
for sh in seq: seq[sh].sort()
ptr={sh:0 for sh in seq}
for r in range(2,ws.max_row+1):
    if not ws.cell(r,4).value: continue
    sh=inv.get(ws.cell(r,1).value)
    if sh not in seq or ptr[sh]>=len(seq[sh]): continue
    v=seq[sh][ptr[sh]][1]; ptr[sh]+=1
    for j,x in enumerate(v):
        cell=ws.cell(r,n0+1+j,x); 
        if j in(2,3): cell.number_format='#,##0.00 €'
ws.auto_filter.ref=f"A1:{get_column_letter(n0+len(NEW))}{ws.max_row}"
# İçindekiler notları
ic=wb['İçindekiler']; r=ic.max_row+2
notes=["Web sitesi sütunları (her ürün sekmesinin sonunda, mavi başlıklı)",
"• Sitede Aktif: Evet ise ürün sitede görünür, Hayır ise gizlenir.",
"• Site Kategorisi: ürünün sitedeki yeri (Ana > Grup > Alt). Geçerli yollar 'Site Kategorileri' sekmesindedir; hücrede açılır listeden seçilebilir.",
"• Piyasa Ortalaması (EUR, KDV hariç): Türkiye'deki Siemens satış sitelerinin fiyat ortalaması. Fiyat Dayanağı 'Tahmini' olan satırlarda, aynı ailedeki ürünlerin piyasa/liste oranıyla liste fiyatından hesaplanmıştır.",
"• Nihai Satış Fiyatı (EUR, KDV hariç): sitede satılacak fiyat. Başlangıçta piyasa ortalamasına eşittir; istediğiniz gibi değiştirebilirsiniz. Sitede TL karşılığı TCMB kuruyla hesaplanır.",
"• Stok Adedi: depodaki adet. Boş ya da 0 ise ürün sitede 'Stokta yok' olarak işaretlenir; sipariş adedi stoğu aşarsa fazlası 'tedarikle' görünür.",
"• Site bu sütunları ürün sekmelerinden okur. 'Tüm Ürünler' sekmesi referans kopyasıdır ve siteye aktarılmaz.",
"• Düzenledikten sonra dosyayı sitedeki 'Katalog yönetimi' bölümünden yükleyin."]
for i,t in enumerate(notes):
    c=ic.cell(r+i,1,t)
    if i==0: c.font=Font(b=True)
wb.save(OUT)
print(stats)
